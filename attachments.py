"""Inbound WeChat attachments: download via Hub, extract text, describe images."""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".yaml", ".yml",
    ".toml", ".ini", ".cfg", ".conf", ".log", ".xml", ".html", ".htm", ".css",
    ".js", ".ts", ".jsx", ".tsx", ".py", ".go", ".java", ".rs", ".c", ".h",
    ".cpp", ".hpp", ".sh", ".bash", ".sql", ".env.example",
}
MAX_TEXT_CHARS = 20000
MAX_FILES = 3
MAX_IMAGE_BYTES = 4 * 1024 * 1024
MAX_FILE_BYTES = 10 * 1024 * 1024


class AttachmentError(Exception):
    pass


class Attachments:
    def __init__(self, hub, llm, base_dir: str = "/workspace/inbox", *,
                 enabled: bool = True, max_files: int = MAX_FILES,
                 max_image_bytes: int = MAX_IMAGE_BYTES,
                 max_file_bytes: int = MAX_FILE_BYTES,
                 max_text_chars: int = MAX_TEXT_CHARS):
        self.hub = hub
        self.llm = llm
        self.enabled = enabled
        self.base_dir = Path(base_dir)
        self.max_files = max(1, int(max_files))
        self.max_image_bytes = max(1024, int(max_image_bytes))
        self.max_file_bytes = max(1024, int(max_file_bytes))
        self.max_text_chars = max(1000, int(max_text_chars))
        self.base_dir.mkdir(parents=True, exist_ok=True)

    # ------------------------------------------------------------------
    def collect(self, message: dict) -> dict:
        """Download attachments and return image data URIs plus text notes."""
        from hub import parse_attachments

        attachments = parse_attachments(message)
        result = {"items": len(attachments) if self.enabled else 0,
                  "images": [], "notes": [], "errors": [], "saved": []}
        if not self.enabled or not attachments:
            return result
        for item in attachments[: self.max_files]:
            kind = item.get("type") or "file"
            name = self._name_for(item)
            limit = self.max_image_bytes if kind == "image" else self.max_file_bytes
            try:
                data, content_type = self.hub.fetch_media(item, limit)
            except Exception as exc:
                result["errors"].append(f"{name} 下载失败：{exc}")
                continue
            folder = self.base_dir / self._folder_name()
            folder.mkdir(parents=True, exist_ok=True)
            path = folder / name
            path.write_bytes(data)
            result["saved"].append(str(path))
            if kind == "image":
                mime = (content_type or "").split(";")[0].strip().lower()
                if not mime.startswith("image/"):
                    mime = self._guess_image_mime(name, data)
                encoded = base64.b64encode(data).decode("ascii")
                result["images"].append(f"data:{mime};base64,{encoded}")
            elif kind == "file":
                text, mode = self._extract_text(data, name)
                if text:
                    clipped = text[: self.max_text_chars]
                    more = "（内容已截断）" if len(text) > self.max_text_chars else ""
                    result["notes"].append(f"文件 {name} 的内容{more}：\n{clipped}")
                else:
                    result["notes"].append(
                        f"文件 {name} 已保存到 {path}，暂不支持解析该格式（{mode}）。"
                    )
            else:
                result["notes"].append(
                    f"{kind} {name} 已保存到 {path}，当前只能保存，不能理解其中的音视频内容。"
                )
        if len(attachments) > self.max_files:
            result["notes"].append(
                f"附件共 {len(attachments)} 个，只处理了前 {self.max_files} 个。"
            )
        return result

    def describe(self, images: list[str], prompt: str) -> str:
        content = [{"type": "text", "text": prompt}]
        content.extend({"type": "image_url", "image_url": {"url": uri}} for uri in images)
        message = self.llm.complete([{"role": "user", "content": content}], None, max_tokens=1200)
        text = (message.get("content") or "").strip()
        if not text:
            raise AttachmentError("视觉模型没有返回内容")
        return text

    # ------------------------------------------------------------------
    def _folder_name(self) -> str:
        import time
        import uuid
        return f"{time.strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:8]}"

    @staticmethod
    def _name_for(item: dict) -> str:
        kind = str(item.get("type") or "file")
        raw = str(item.get("file_name") or "").strip()
        name = Path(raw).name if raw else ""
        if not name:
            name = {"image": "image.jpg", "video": "video.mp4", "voice": "voice.wav"}.get(
                kind, "file.bin"
            )
        name = re.sub(r"[^\w.\-()\u4e00-\u9fff ]+", "_", name).strip()
        return (name or "file.bin")[:80]

    @staticmethod
    def _guess_image_mime(name: str, data: bytes) -> str:
        suffix = Path(name).suffix.lower()
        if suffix in {".png"}:
            return "image/png"
        if suffix in {".gif"}:
            return "image/gif"
        if suffix in {".webp"}:
            return "image/webp"
        if data.startswith(b"\x89PNG"):
            return "image/png"
        if data.startswith(b"GIF8"):
            return "image/gif"
        if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
            return "image/webp"
        return "image/jpeg"

    def _extract_text(self, data: bytes, name: str) -> tuple[str, str]:
        suffix = Path(name).suffix.lower()
        if suffix in TEXT_SUFFIXES or suffix == "":
            for encoding in ("utf-8-sig", "utf-8", "gb18030", "utf-16"):
                try:
                    return data.decode(encoding), encoding
                except UnicodeDecodeError:
                    continue
            return "", "binary"
        if suffix == ".pdf":
            try:
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(data))
                pages = [(page.extract_text() or "") for page in reader.pages]
                return "\n".join(pages).strip(), "pdf"
            except Exception as exc:
                return "", f"pdf error: {exc}"
        if suffix == ".docx":
            try:
                import docx
                document = docx.Document(io.BytesIO(data))
                parts = [p.text for p in document.paragraphs if p.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        parts.append("\t".join(cell.text for cell in row.cells))
                return "\n".join(parts).strip(), "docx"
            except Exception as exc:
                return "", f"docx error: {exc}"
        if suffix in {".xlsx", ".xls"}:
            return "", "xlsx-not-installed"
        return "", "unsupported"
